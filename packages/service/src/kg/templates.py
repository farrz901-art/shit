"""Excel / JSON 模板生成

v1 决策: 硬编码 1 个 QA 模板 (列头固定), 后续 PR 升级为从 yaml 动态生成
"""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


def build_qa_excel() -> io.BytesIO:
    """生成 QA 对录入 Excel 模板 (含列头 + 示例行 + 表头样式)

    列头: question, answer, tags (逗号分隔)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None  # 新建 Workbook 后 active 必有
    ws.title = "QA 对"

    headers = ["question", "answer", "tags"]
    ws.append(headers)

    # 表头样式: 蓝底白字
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="3B82F6")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 列宽
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 30

    # 示例行
    ws.append(["1LE0001 是什么系列?", "1LE 系列是西门子低压电机", "电机, 西门子"])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
